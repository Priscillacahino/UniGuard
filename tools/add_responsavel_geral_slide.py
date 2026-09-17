from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path

pptx_path = Path('UniGuard_Apresentacao_Profissional_Atualizada.pptx')
prs = Presentation(pptx_path)

# Evita duplicar a página caso o workflow seja executado mais de uma vez.
for slide in prs.slides:
    for shape in slide.shapes:
        if hasattr(shape, 'text') and 'Responsável geral pelo projeto UniGuard' in shape.text:
            prs.save(pptx_path)
            raise SystemExit('Slide de responsabilidade geral já existe.')

blank = prs.slide_layouts[6]
slide = prs.slides.add_slide(blank)

navy = RGBColor(11, 46, 77)
cyan = RGBColor(18, 181, 203)
white = RGBColor(255, 255, 255)
soft = RGBColor(223, 233, 241)
muted = RGBColor(173, 188, 202)

slide.background.fill.solid()
slide.background.fill.fore_color.rgb = navy

shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.15), Inches(0), Inches(4.2), Inches(7.5))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(7, 37, 62)
shape.line.fill.background()

shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(0.65), Inches(1.65), Inches(0.05))
shape.fill.solid()
shape.fill.fore_color.rgb = cyan
shape.line.fill.background()

shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.55), Inches(1.1), Inches(2.9), Inches(2.9))
shape.fill.solid()
shape.fill.fore_color.rgb = RGBColor(7, 58, 92)
shape.line.fill.background()

shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.05), Inches(1.6), Inches(1.9), Inches(1.9))
shape.fill.solid()
shape.fill.fore_color.rgb = cyan
shape.line.fill.background()


def add_text(x, y, w, h, text, size, bold=False, color=white, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = 'Aptos'
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return tb

add_text(0.65, 0.86, 5.5, 0.35, 'AUTORIA E RESPONSABILIDADE GERAL', 10, True, cyan)
add_text(0.65, 1.45, 7.6, 0.8, 'Priscilla Cahino', 34, True, white)
add_text(0.65, 2.18, 7.9, 0.5, 'Responsável geral pelo projeto UniGuard', 17, True, RGBColor(218, 239, 246))
add_text(
    0.65,
    3.05,
    7.65,
    1.2,
    'Atuação responsável pela idealização, concepção do produto, definição da proposta de valor, estruturação dos fluxos, requisitos, regras de negócio, prototipação, documentação, organização dos artefatos e consolidação da apresentação profissional do projeto.',
    14,
    False,
    soft,
)

card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(4.75), Inches(7.75), Inches(1.2))
card.fill.solid()
card.fill.fore_color.rgb = RGBColor(15, 65, 104)
card.line.color.rgb = RGBColor(38, 98, 136)
add_text(
    0.95,
    5.02,
    7.15,
    0.62,
    'Projeto acadêmico e tecnológico em evolução, com finalidade demonstrativa e base documental alinhada ao repositório GitHub.',
    13,
    False,
    RGBColor(229, 240, 246),
)

add_text(9.65, 4.3, 2.5, 0.35, 'UNIGUARD', 11, True, cyan, PP_ALIGN.CENTER)
add_text(9.35, 4.75, 3.1, 0.8, 'Segurança universitária\nconectada ao cuidado', 18, True, white, PP_ALIGN.CENTER)
add_text(0.65, 6.92, 5.3, 0.25, 'PROTÓTIPO ACADÊMICO • DADOS DEMONSTRATIVOS', 8, False, muted)
add_text(11.72, 6.88, 0.45, 0.25, str(len(prs.slides)).zfill(2), 8, False, muted, PP_ALIGN.RIGHT)

prs.save(pptx_path)
print(f'Apresentação atualizada com {len(prs.slides)} slides.')
