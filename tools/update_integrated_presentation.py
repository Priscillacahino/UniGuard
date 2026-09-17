from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path

pptx_path = Path('UniGuard_Apresentacao_Profissional_Atualizada.pptx')
prs = Presentation(pptx_path)

def replace(slide, old, new):
    for sh in slide.shapes:
        if hasattr(sh, 'text') and sh.text.strip() == old:
            sh.text = new
            return

for old,new in {
    'Cadastro, contexto e consentimento':'Cadastro, verificação e vínculo',
    'Autenticação e registro seguro':'Autenticação, regras e auditoria',
    'Triagem, despacho e conclusão':'Triagem, status e encerramento'
}.items(): replace(prs.slides[2], old, new)

s=prs.slides[4]
for old,new in {
'Cadastro inclusivo e preparado para o contexto':'Cadastro com identidade verificada e contexto institucional',
'A validação foi ampliada para aceitar qualquer e-mail válido — inclusive visitantes em trânsito no campus.':'O cadastro utiliza e-mail e senha, mas o uso institucional exige confirmação do e-mail e vínculo ativo com a instituição autorizada.',
'Acesso mais inclusivo':'Acesso controlado',
'Qualquer pessoa pode iniciar o fluxo sem depender de domínio institucional.':'A credencial não basta: o usuário precisa ter e-mail verificado e membership institucional ativa.',
'Identificação contextual':'Identificação institucional',
'Vínculo, matrícula/CPF, telefone e contato de emergência apoiam a triagem.':'O institutionId separa usuários, operadores e alertas entre instituições, evitando acesso cruzado.',
'Coletar somente o necessário, com consentimento e finalidade explícita.':'Dados médicos e medicalNotes foram retirados/bloqueados; a coleta permanece limitada ao necessário para a ocorrência.'
}.items(): replace(s,old,new)

s=prs.slides[6]
replace(s,'Authentication • Firestore','Authentication • Firestore • Storage')
replace(s,'O mesmo Firestore conecta a experiência do usuário à visão do operador.','Firebase conecta usuário e central com regras por institutionId, trilha de auditoria e evidências protegidas.')

s=prs.slides[7]
for old,new in {
'Autenticação Firebase':'Autenticação Firebase por e-mail/senha',
'Regras de acesso no Firestore':'E-mail verificado + vínculo institucional ativo',
'Perfis separados: usuário e operador':'Isolamento por institutionId e perfis separados',
'Bloqueio de exclusão de alertas':'Auditoria append-only + evidências sem sobrescrita',
'Verificação de propriedade do e-mail':'Ativar TTL remoto no Firestore',
'Política de privacidade e consentimento':'Configurar ciclo de vida do Storage',
'Retenção e anonimização de dados':'Habilitar App Check e rate limit confiável',
'App Check, logs e auditoria institucional':'Validar fluxo ponta a ponta e testes negativos'
}.items(): replace(s,old,new)

s=prs.slides[8]
for old,new in {
'Validação técnica do fluxo principal':'Validação técnica e pontos ainda pendentes',
'Testes realizados no ambiente de desenvolvimento e na implantação demonstrativa.':'O repositório contém testes e controles implementados, mas a comprovação integrada Android → Firebase → Central ainda precisa ser concluída e registrada.',
'Cadastro e login com e-mail válido':'Cadastro, login e verificação de e-mail',
'Criação e leitura de alertas':'Regras, auditoria e isolamento institucional',
'Envio do alerta com localização':'SOS com limites locais contra abuso',
'Visualização e mudança de status':'Recebimento, visualização e mudança de status',
'GitHub conectado à Vercel':'Deploy web e CI disponíveis',
'Fluxo acessível pelo navegador do celular':'Modo demonstrativo local disponível',
'MVP funcional: integração técnica do fluxo principal concluída.':'Status correto: protótipo acadêmico avançado e MVP demonstrativo em desenvolvimento, ainda dependente de validação integrada.'
}.items(): replace(s,old,new)

s=prs.slides[10]
for old,new in {
'MVP conectado':'MVP demonstrativo',
'Fluxo SOS + Firebase + painel':'Fluxo SOS + controles + painel',
'Pesquisa e piloto':'Validação integrada',
'Questionário, cenários e testes com usuários':'Testes Android → Firebase → Central, isolamento e abuso',
'LGPD, protocolos e integrações':'TTL, Storage lifecycle, App Check e protocolos',
'Operação institucional':'Piloto institucional',
'Monitoramento, SLAs e análise preventiva':'Autorizações, SLAs, operação e análise preventiva',
'Próxima validação: compreender a percepção de segurança da comunidade e ajustar a solução antes de uma implantação institucional.':'Próxima validação: comprovar o fluxo ponta a ponta com contas de teste, documentar evidências e só então avançar para um piloto institucional.'
}.items(): replace(s,old,new)

blank=prs.slide_layouts[6]
navy=RGBColor(11,46,77); cyan=RGBColor(18,181,203); slate=RGBColor(100,110,120)
light=RGBColor(244,248,251); white=RGBColor(255,255,255); green=RGBColor(16,158,120); blue=RGBColor(0,120,181)

def box(slide,x,y,w,h,text,size=18,bold=False,color=navy,align=PP_ALIGN.LEFT):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=text; r.font.name='Aptos'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color

def req_slide(kicker,title,left_h,left,right_h,right,num):
    s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=light
    q=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.6), Inches(.45), Inches(1.4), Inches(.04)); q.fill.solid(); q.fill.fore_color.rgb=cyan; q.line.fill.background()
    box(s,.6,.58,2.8,.25,kicker.upper(),10,True,cyan); box(s,.6,.95,11.3,.6,title,27,True,navy)
    for x,head,lines,accent in [(.6,left_h,left,blue),(6.55,right_h,right,green)]:
        card=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.8), Inches(5.35), Inches(4.65)); card.fill.solid(); card.fill.fore_color.rgb=white; card.line.color.rgb=RGBColor(222,229,235)
        bar=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(1.8), Inches(5.35), Inches(.07)); bar.fill.solid(); bar.fill.fore_color.rgb=accent; bar.line.fill.background()
        box(s,x+.25,2.05,4.85,.35,head,15,True,navy); yy=2.55
        for line in lines: box(s,x+.28,yy,4.78,.53,'• '+line,11,False,RGBColor(34,40,49)); yy+=.58
    hr=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.6), Inches(7.05), Inches(11.25), Inches(.012)); hr.fill.solid(); hr.fill.fore_color.rgb=RGBColor(220,226,232); hr.line.fill.background()
    box(s,.6,7.12,5.5,.2,'PROTÓTIPO ACADÊMICO • DADOS DEMONSTRATIVOS',8,False,slate); box(s,11.42,7.08,.4,.25,f'{num:02d}',8,False,slate,PP_ALIGN.RIGHT)

req_slide('REQUISITOS','Requisitos funcionais aplicados','RF01–RF06',[
'RF01 – Cadastrar e autenticar por e-mail e senha.','RF02 – Exigir e-mail verificado e vínculo institucional ativo.','RF03 – Associar usuário à instituição/campus autorizado.','RF04 – Obter localização mediante permissão.','RF05 – Criar SOS com categoria, data/hora e localização.','RF06 – Anexar evidência fotográfica quando autorizada.'],
'RF07–RF12',['RF07 – Exibir alertas ao operador da mesma instituição.','RF08 – Registrar recebimento, visualização e mudanças de status.','RF09 – Permitir acompanhamento e histórico do alerta.','RF10 – Registrar eventos de auditoria.','RF11 – Bloquear criação abusiva/repetitiva de alertas.','RF12 – Limpar dados temporários locais conforme retenção.'],13)

req_slide('REQUISITOS','Requisitos não funcionais aplicados','RNF01–RNF05',[
'RNF01 – Acesso controlado por perfil, conta e instituição.','RNF02 – Minimização de dados; sem prontuário médico.','RNF03 – Auditoria append-only e rastreável.','RNF04 – Retenção prevista: alertas 90 dias; auditoria 365 dias.','RNF05 – Evidências 30 dias; temporários locais 7 dias.'],
'RNF06–RNF10',['RNF06 – Evidência não pode ser sobrescrita pelo cliente.','RNF07 – Nenhum acesso cruzado entre instituições.','RNF08 – Modo demonstrativo local independente do Firebase.','RNF09 – Permissões necessárias e confirmação clara ao usuário.','RNF10 – MySQL futuro somente via API/backend seguro.'],14)

req_slide('NEGÓCIO','Regras de negócio aplicadas','RN01–RN06',[
'RN01 – E-mail deve estar verificado para operar.','RN02 – Membership ativa e institutionId são obrigatórios.','RN03 – Operador atua somente na própria instituição.','RN04 – Máximo de 1 alerta ativo por usuário.','RN05 – Máximo de 3 alertas em 30 minutos.','RN06 – Intervalo mínimo de 60 segundos entre novos alertas.'],
'RN07–RN12',['RN07 – Evidência enviada não pode ser substituída.','RN08 – medicalNotes e dados clínicos são proibidos.','RN09 – Auditoria registra ator, alerta, instituição e horário.','RN10 – Exclusão antecipada exige processo administrativo auditável.','RN11 – Legado sem institutionId/retentionExpiresAt deve ser migrado ou isolado.','RN12 – Em validação pública, utilizar somente dados fictícios.'],15)

prs.save(pptx_path)
